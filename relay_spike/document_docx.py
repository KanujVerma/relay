"""Bounded, bytes-only DOCX inspection for the synthetic import slice."""
from __future__ import annotations

from io import BytesIO
from pathlib import PurePosixPath
from urllib.parse import urlsplit
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree as SafeET


MAX_DOCX_ENTRIES = 256
MAX_DOCX_UNCOMPRESSED_BYTES = 8 * 1024 * 1024
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAIN_DOCUMENT_RELATIONSHIP = OFFICE_REL_NS + "/officeDocument"
_PROPERTY_CONTENT_TAGS = frozenset({
    f"{{{WORD_NS}}}p",
    f"{{{WORD_NS}}}r",
    f"{{{WORD_NS}}}hyperlink",
    f"{{{WORD_NS}}}t",
    f"{{{WORD_NS}}}delText",
    f"{{{WORD_NS}}}instrText",
})
_UNSUPPORTED_WORD_TAGS = frozenset({
    "tbl", "drawing", "pict", "object", "altChunk", "hyperlink",
    "footnoteReference", "endnoteReference", "commentRangeStart",
    "commentRangeEnd", "fldSimple", "instrText", "delText", "ins",
    "del", "moveFrom", "moveTo",
})


class InvalidDocx(ValueError):
    """Malformed or over-limit archive, with a non-content reason code."""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


class UnsupportedDocx(ValueError):
    """Valid DOCX containing structures this slice cannot locate exactly."""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


def _require_leaf(node, *, allow_text: bool = False) -> None:
    """Enforce the accepted Word leaf grammar before projecting its content."""
    if len(node) or (not allow_text and (node.text or "").strip()):
        raise UnsupportedDocx("docx_word_leaf_invalid")


def _text_free_properties(node) -> None:
    """Reject content hidden in property containers python-docx skips."""
    for item in node.iter():
        if item is not node and item.tag in _PROPERTY_CONTENT_TAGS:
            raise UnsupportedDocx("docx_unsupported_property_content")
        if (item.text or "").strip() or (item.tail or "").strip():
            raise UnsupportedDocx("docx_text_outside_supported_run")


def _run_text(run) -> str:
    if (run.text or "").strip():
        raise UnsupportedDocx("docx_text_outside_supported_run")
    parts: list[str] = []
    seen_content = False
    seen_properties = False
    for child in run:
        if child.tag == f"{{{WORD_NS}}}rPr":
            if seen_properties or seen_content:
                raise UnsupportedDocx("docx_run_structure_unsupported")
            seen_properties = True
            _text_free_properties(child)
        elif child.tag == f"{{{WORD_NS}}}t":
            _require_leaf(child, allow_text=True)
            seen_content = True
            parts.append(child.text or "")
        elif child.tag == f"{{{WORD_NS}}}tab":
            _require_leaf(child)
            seen_content = True
            parts.append("\t")
        elif child.tag == f"{{{WORD_NS}}}br":
            _require_leaf(child)
            seen_content = True
            break_type = child.attrib.get(f"{{{WORD_NS}}}type", "textWrapping")
            if break_type not in {"textWrapping", "page", "column"}:
                raise UnsupportedDocx("docx_break_type_unsupported")
            parts.append("\n" if break_type == "textWrapping" else "")
        elif child.tag == f"{{{WORD_NS}}}cr":
            _require_leaf(child)
            seen_content = True
            parts.append("\n")
        elif child.tag == f"{{{WORD_NS}}}noBreakHyphen":
            _require_leaf(child)
            seen_content = True
            parts.append("-")
        elif child.tag == f"{{{WORD_NS}}}ptab":
            _require_leaf(child)
            seen_content = True
            parts.append("\t")
        else:
            # Structural wrappers, drawings, fields, revision marks, and future
            # text-bearing nodes are not part of python-docx's Run.text result.
            raise UnsupportedDocx("docx_unsupported_inline_content")
        if (child.tail or "").strip():
            raise UnsupportedDocx("docx_text_outside_supported_run")
    return "".join(parts)


def _source_paragraph_texts(document_xml: bytes) -> list[str]:
    """Validate the complete body shape and project supported source text."""
    try:
        root = SafeET.fromstring(document_xml)
    except Exception:
        raise InvalidDocx("docx_document_xml_invalid") from None
    if root.tag != f"{{{WORD_NS}}}document" or len(root) != 1:
        raise UnsupportedDocx("docx_document_structure_unsupported")
    body = root[0]
    if body.tag != f"{{{WORD_NS}}}body":
        raise UnsupportedDocx("docx_document_structure_unsupported")
    if ((root.text or "").strip() or (root.tail or "").strip() or
            (body.text or "").strip() or (body.tail or "").strip()):
        raise UnsupportedDocx("docx_text_outside_supported_body")

    paragraphs: list[str] = []
    for child in body:
        if child.tag == f"{{{WORD_NS}}}p":
            if (child.text or "").strip():
                raise UnsupportedDocx("docx_text_outside_supported_paragraph")
            paragraph_text: list[str] = []
            seen_properties = False
            seen_run = False
            for item in child:
                if item.tag == f"{{{WORD_NS}}}pPr":
                    if seen_properties or seen_run:
                        raise UnsupportedDocx("docx_paragraph_structure_unsupported")
                    seen_properties = True
                    _text_free_properties(item)
                elif item.tag == f"{{{WORD_NS}}}r":
                    seen_run = True
                    paragraph_text.append(_run_text(item))
                else:
                    raise UnsupportedDocx("docx_unsupported_inline_content")
                if (item.tail or "").strip():
                    raise UnsupportedDocx("docx_text_outside_supported_paragraph")
            paragraphs.append("".join(paragraph_text))
        elif child.tag == f"{{{WORD_NS}}}sectPr":
            _text_free_properties(child)
        else:
            # Body wrappers such as w:sdt, w:customXml, mc:AlternateContent,
            # tables, and text boxes can contain paragraphs omitted by
            # python-docx's Document.paragraphs collection.
            raise UnsupportedDocx("docx_unsupported_body_content")
        if (child.tail or "").strip():
            raise UnsupportedDocx("docx_text_outside_supported_body")
    return paragraphs


def _root_relationships(archive: ZipFile, names: set[str]):
    """Require one canonical main-part target matching python-docx behavior."""
    if "_rels/.rels" not in names:
        raise InvalidDocx("docx_package_relationships_missing")
    try:
        relationships = SafeET.fromstring(archive.read("_rels/.rels"))
    except Exception:
        raise InvalidDocx("docx_relationships_invalid") from None
    if relationships.tag != f"{{{REL_NS}}}Relationships":
        raise InvalidDocx("docx_relationships_invalid")

    ids: set[str] = set()
    main_relations = []
    for relation in relationships.findall(f"{{{REL_NS}}}Relationship"):
        rel_id = relation.attrib.get("Id", "")
        if not rel_id or rel_id in ids:
            raise UnsupportedDocx("docx_package_relationship_ambiguous")
        ids.add(rel_id)
        if relation.attrib.get("Type") == MAIN_DOCUMENT_RELATIONSHIP:
            main_relations.append(relation)
    if len(main_relations) != 1:
        raise UnsupportedDocx("docx_main_document_relationship_ambiguous")

    relation = main_relations[0]
    target_mode = relation.attrib.get("TargetMode", "").strip().casefold()
    target = relation.attrib.get("Target", "")
    parsed_target = urlsplit(target)
    if (target_mode not in {"", "internal"} or target != "word/document.xml" or
            parsed_target.scheme or parsed_target.netloc or target not in names):
        raise UnsupportedDocx("docx_main_document_target_unsupported")
    return relationships


def inspect_docx(data: bytes):
    """Return python-docx top-level paragraphs after fail-closed preflight.

    No path is opened and no relationship target is followed. Only plain body
    paragraphs are supported; tables, embedded objects, external relations,
    and non-empty header/footer/note content are rejected.
    """
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise InvalidDocx("docx_entry_limit")

            names: set[str] = set()
            total_size = 0
            for entry in entries:
                name = entry.filename.replace("\\", "/")
                path = PurePosixPath(name)
                if (not name or name.startswith("/") or ":" in name or
                        ".." in path.parts or name in names):
                    raise InvalidDocx("docx_archive_path")
                names.add(name)
                if entry.flag_bits & 0x1:
                    raise UnsupportedDocx("docx_encrypted_member")
                mode = (entry.external_attr >> 16) & 0o170000
                if mode == 0o120000:
                    raise UnsupportedDocx("docx_link_member")
                if entry.file_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise InvalidDocx("docx_member_size_limit")
                total_size += entry.file_size
                if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise InvalidDocx("docx_uncompressed_size_limit")
                if name.startswith("word/embeddings/") or name.startswith("word/media/"):
                    raise UnsupportedDocx("docx_embedded_content")

            if "word/document.xml" not in names:
                raise InvalidDocx("docx_document_missing")

            root_relationships = _root_relationships(archive, names)
            for name in names:
                if not name.endswith(".rels"):
                    continue
                if name == "_rels/.rels":
                    relationships = root_relationships
                else:
                    try:
                        relationships = SafeET.fromstring(archive.read(name))
                    except Exception:
                        raise InvalidDocx("docx_relationships_invalid") from None
                for relationship in relationships.iter(f"{{{REL_NS}}}Relationship"):
                    target_mode = relationship.attrib.get("TargetMode", "").strip().casefold()
                    parsed_target = urlsplit(relationship.attrib.get("Target", ""))
                    if (target_mode == "external" or
                            bool(parsed_target.scheme or parsed_target.netloc)):
                        raise UnsupportedDocx("docx_external_relationship")

            try:
                document_xml = archive.read("word/document.xml")
            except Exception:
                raise InvalidDocx("docx_document_xml_invalid") from None
            source_texts = _source_paragraph_texts(document_xml)

            # The plain-paragraph path does not read these content layers.
            for name in names:
                is_secondary_text = (
                    name.startswith("word/header") or
                    name.startswith("word/footer") or
                    name in {"word/comments.xml", "word/footnotes.xml", "word/endnotes.xml"}
                )
                if not is_secondary_text or not name.endswith(".xml"):
                    continue
                try:
                    root = SafeET.fromstring(archive.read(name))
                except Exception:
                    raise InvalidDocx("docx_secondary_xml_invalid") from None
                for node in root.iter():
                    if not isinstance(node.tag, str) or not node.tag.startswith("{"):
                        continue
                    namespace, local_name = node.tag[1:].split("}", 1)
                    if namespace != WORD_NS:
                        continue
                    if local_name in _UNSUPPORTED_WORD_TAGS:
                        raise UnsupportedDocx("docx_secondary_structure")
                    if local_name == "t" and (node.text or ""):
                        raise UnsupportedDocx("docx_secondary_text")

        from docx import Document

        paragraphs = Document(BytesIO(data)).paragraphs
        if len(paragraphs) != len(source_texts):
            raise UnsupportedDocx("docx_paragraph_coverage_mismatch")
        if any(source != paragraph.text for source, paragraph in zip(source_texts, paragraphs)):
            raise UnsupportedDocx("docx_text_coverage_mismatch")
        return paragraphs
    except (InvalidDocx, UnsupportedDocx):
        raise
    except ImportError:
        raise
    except BadZipFile:
        raise InvalidDocx("docx_archive_invalid") from None
    except Exception:
        raise InvalidDocx("docx_parse_failed") from None
